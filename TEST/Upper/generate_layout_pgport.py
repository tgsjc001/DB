import argparse
import psycopg2
import json
import ctypes
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---------------------------
# Color definitions
# ---------------------------
COLOR_INDIGO = "4B0082"
COLOR_YELLOW = "FFFF00"
COLOR_LIME = "32CD32"

# ---------------------------
# Template and Output paths
# ---------------------------
TEMPLATE_UPPER = "Master Layout Upper.docx"
TEMPLATE_LOWER = "Master Layout Lower.docx"
OUTPUT_UPPER = "Layout_Upper_Generated.docx"
OUTPUT_LOWER = "Layout_Lower_Generated.docx"


# ---------------------------
# Helper Functions
# ---------------------------

def set_cell_background(cell, color_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def parse_table_numbers(table_str):
    table_list = []
    parts = [p.strip() for p in table_str.split(',') if p.strip()]
    for part in parts:
        if '-' in part:
            start, end = part.split('-')
            start, end = start.strip(), end.strip()
            if start.isdigit() and end.isdigit():
                for i in range(int(start), int(end) + 1):
                    table_list.append(str(i))
            elif re.match(r'^\d+[A-Z]$', start) and re.match(r'^\d+[A-Z]$', end):
                base = re.findall(r'^\d+', start)[0]
                start_letter = start[-1]
                end_letter = end[-1]
                for c in range(ord(start_letter), ord(end_letter) + 1):
                    table_list.append(f"{base}{chr(c)}")
            else:
                table_list.append(start)
                table_list.append(end)
        else:
            table_list.append(part)
    return table_list


def build_tags(v):
    counts = {
        "G": int(v.get("six_ft_gun") or 0) + int(v.get("eight_ft_gun") or 0),
        "K": int(v.get("six_ft_knife") or 0) + int(v.get("eight_ft_knife") or 0),
        "D": int(v.get("six_ft_display") or 0) + int(v.get("eight_ft_display") or 0),
        "N": int(v.get("six_ft_nonrelated") or 0) + int(v.get("eight_ft_nonrelated") or 0),
    }
    main_tag = max(counts, key=counts.get) if any(counts.values()) else ""

    # 👇 Add R if any rollover flag is true, regardless of per-table count
    if v.get("on_rollover") or v.get("rta") or v.get("rtn"):
        return f"{main_tag}R" if main_tag else "R"

    return main_tag


def load_vendors_from_postgres(host, port, dbname, user, password, table="vendors"):
    conn = psycopg2.connect(host=host, port=port, dbname=dbname, user=user, password=password)
    cur = conn.cursor()
    cur.execute(f"""
        SELECT first_name, last_name, business_name, island, row, table_numbers,
               six_ft_gun, six_ft_knife, six_ft_display, six_ft_nonrelated,
               eight_ft_gun, eight_ft_knife, eight_ft_display, eight_ft_nonrelated,
               electricity, on_rollover, rta, rtn, saturday_setup
        FROM {table};
    """)
    columns = [desc[0] for desc in cur.description]
    vendors = [dict(zip(columns, row)) for row in cur.fetchall()]
    cur.close()
    conn.close()
    return vendors


def dump_vendor_debug(vendors, filename="vendor_debug.txt"):
    try:
        with open(filename, "w", encoding="utf-8") as f:
            for i, v in enumerate(vendors, start=1):
                f.write(f"Vendor {i}:\n")
                f.write(json.dumps(v, indent=4, default=str))
                f.write("\n" + "-" * 80 + "\n")
        print(f"[DEBUG] Vendor data dump written to {filename}")
    except Exception as e:
        print(f"[ERROR] Failed to write vendor debug dump: {e}")


def check_double_bookings(vendors, filename="double_booking_report.txt"):
    conflicts = []
    seen = {}

    for v in vendors:
        island = (v.get("island") or "").strip().upper()
        row = (v.get("row") or "").strip().upper()
        name = v.get("business_name") or f"{v.get('first_name', '')} {v.get('last_name', '')}".strip()
        tables = parse_table_numbers(v.get("table_numbers"))

        for t in tables:
            key = (island, row, t)
            if key in seen:
                conflicts.append(f"{island} {row} Table {t} -> {seen[key]} & {name}")
            else:
                seen[key] = name

    if conflicts:
        with open(filename, "w", encoding="utf-8") as f:
            f.write("Double Booking Report\n")
            f.write("====================\n\n")
            for c in conflicts:
                f.write(c + "\n")

        conflict_msg = f"{len(conflicts)} double bookings detected! See {filename}"
        print(f"[WARNING] {conflict_msg}")
        try:
            ctypes.windll.user32.MessageBoxW(0, conflict_msg, "Double Booking Alert", 0x10)
        except Exception:
            print("[WARN] Unable to show popup alert.")
        raise SystemExit(f"Aborting due to double bookings. Check {filename} for details.")
    else:
        print("[INFO] No double bookings detected.")


# ---------------------------
# Document Filling
# ---------------------------

def fill_document(template_path, vendors, output_path):
    from docx import Document
    doc = Document(template_path)
    vendor_map = {}
    skipped_vendors = []

    with open("vendor_debug.txt", "w", encoding="utf-8") as f:
        for i, v in enumerate(vendors, 1):
            f.write(f"Vendor {i}:\n")
            f.write(json.dumps(v, indent=4, default=str))
            f.write("\n" + "-" * 80 + "\n")
    print("[DEBUG] Vendor data dump written to vendor_debug.txt")

    for v in vendors:
        name = (v.get("business_name") or f"{v.get("first_name", "")} {v.get("last_name", "")}").strip() or "Unknown Vendor"
        island = (v.get("island") or "").strip().upper()
        row = (v.get("row") or "").strip().upper()
        tables = parse_table_numbers(v.get("table_numbers") or "")
        if not island or not row or not tables:
            skipped_vendors.append(name)
            continue
        key = f"{island}{row}"
        for t in tables:
            vendor_map.setdefault(key, {})[str(t)] = {
                "name": name,
                "all_tables": tables,
                "tags": build_tags(v),
                "electric": v.get("electricity", False),
                "rollover": v.get("rta", False) or v.get("rtn", False) or v.get("on_rollover", False),
                "saturday": v.get("saturday_setup", False)
            }

    filled = []
    unmatched = []
    layout_log = []

    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            try:
                layout_log.append(f"[Table {table_idx}] Row {row_idx}: " + " | ".join(c.text.strip() for c in row.cells))
            except Exception as e:
                layout_log.append(f"[Table {table_idx}] Row {row_idx}: ERROR reading row - {e}")
            for section_col, table_col, name_col, tag_col in [(1,2,3,4), (6,7,8,9)]:
                if len(row.cells) <= tag_col:
                    continue
                section = row.cells[section_col].text.strip().upper()
                table_num = row.cells[table_col].text.strip()
                if not section or not table_num:
                    continue
                info = vendor_map.get(section, {}).get(table_num)
                if info:
                    idx = info["all_tables"].index(table_num) if table_num in info["all_tables"] else 0
                    is_seq = all(str(int(info["all_tables"][i])+1) == str(info["all_tables"][i+1]) for i in range(len(info["all_tables"])-1)) if len(info["all_tables"]) > 1 else False
                    display_name = info["name"] if idx == 0 else ("↓" if is_seq else f"{info['name']} (with {','.join(t for t in info['all_tables'] if t != table_num)})")
                    row.cells[name_col].text = display_name
                    row.cells[tag_col].text = info["tags"]
                    for p in row.cells[name_col].paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for p in row.cells[tag_col].paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if info["electric"]:
                        set_cell_background(row.cells[name_col], COLOR_INDIGO)
                        set_cell_background(row.cells[tag_col], COLOR_INDIGO)
                    elif info["rollover"]:
                        set_cell_background(row.cells[name_col], COLOR_YELLOW)
                        set_cell_background(row.cells[tag_col], COLOR_YELLOW)
                    elif info["saturday"]:
                        set_cell_background(row.cells[name_col], COLOR_LIME)
                        set_cell_background(row.cells[tag_col], COLOR_LIME)
                    filled.append(f"{section} Table {table_num} => {display_name} [{info['tags']}]")
                else:
                    unmatched.append(f"{section} Table {table_num}")

    with open("layout_scan_debug.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(layout_log))
    with open("filled_tables_debug.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(filled))
    with open("unmatched_sections_debug.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(unmatched))

    doc.save(output_path)
    print(f"[INFO] Layout saved: {output_path} | {len(filled)} tables filled | {len(unmatched)} unmatched")

def generate_layouts(vendors, template_choice):
    if template_choice in ("upper", "both"):
        fill_document(TEMPLATE_UPPER, vendors, OUTPUT_UPPER)
    if template_choice in ("lower", "both"):
        fill_document(TEMPLATE_LOWER, vendors, OUTPUT_LOWER)


# ---------------------------
# Main Entry Point
# ---------------------------

def main():
    parser = argparse.ArgumentParser(description="Generate Upper & Lower layout documents from vendor data.")
    parser.add_argument("--template", choices=["upper", "lower", "both"], default="both")
    parser.add_argument("--pg_host", default="localhost")
    parser.add_argument("--pg_port", type=int, default=5432)
    parser.add_argument("--pg_dbname", default="your_db")
    parser.add_argument("--pg_user", default="your_user")
    parser.add_argument("--pg_password", default="your_password")
    parser.add_argument("--table", default="vendors")
    args = parser.parse_args()

    print(f"[INFO] Generating layouts using template: {args.template}")
    vendors = load_vendors_from_postgres(
        args.pg_host, args.pg_port, args.pg_dbname, args.pg_user, args.pg_password, args.table
    )
    dump_vendor_debug(vendors)
    check_double_bookings(vendors)
    generate_layouts(vendors, args.template)


if __name__ == "__main__":
    main()
