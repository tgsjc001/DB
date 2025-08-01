import argparse
import psycopg2
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import re

COLOR_INDIGO = "4B0082"
COLOR_YELLOW = "FFFF00"
COLOR_LIME = "32CD32"

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.shared import Pt


def center_cell_text(cell):
	for paragraph in cell.paragraphs:
		paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


ROWS_WITH_ISLAND = ['LA', 'LB', 'LC', 'LD', 'XLN']
ROWS_NO_ISLAND = ['LN', 'LNW', 'LS', 'LSW', 'LE', 'LW', 'XLS']


def set_cell_background(cell, color_hex):
	shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
	cell._tc.get_or_add_tcPr().append(shading)


def parse_table_numbers(table_str):
	table_list = []
	parts = [p.strip() for p in table_str.split(',') if p.strip()]
	for part in parts:
		if '-' in part:
			start, end = part.split('-')
			start, end = start.strip(), end.strip()
			if start.isdigit() and end.isdigit():
				for i in range(int(start), int(end) + 1):
					table_list.append(str(i))
			elif re.match(r'^\d+[A-Z]$', start) and re.match(r'^\d+[A-Z]$', end):
				base = re.findall(r'^\d+', start)[0]
				start_letter = start[-1]
				end_letter = end[-1]
				for c in range(ord(start_letter), ord(end_letter) + 1):
					table_list.append(f"{base}{chr(c)}")
			else:
				table_list.append(start)
				table_list.append(end)
		else:
			table_list.append(part)
	return table_list


def load_vendors_from_postgres(host, port, dbname, user, password):
	conn = psycopg2.connect(host=host, port=port, dbname=dbname, user=user, password=password)
	cur = conn.cursor()
	cur.execute("""
		SELECT first_name, last_name, business_name, island, row, table_numbers,
			electricity, on_rollover, rta, rtn, saturday_setup, taca_member,
			six_ft_gun, eight_ft_gun, six_ft_knife, eight_ft_knife,
			six_ft_display, eight_ft_display, six_ft_nonrelated, eight_ft_nonrelated
		FROM vendors
	""")
	columns = [desc[0] for desc in cur.description]
	vendors = [dict(zip(columns, row)) for row in cur.fetchall()]
	cur.close()
	conn.close()

	with open("vendor_debug.txt", "w", encoding="utf-8") as f:
		for v in vendors:
			f.write(str(v) + "\n")
	return vendors


def build_tags(v):
	counts = {
		"G": int(v.get("six_ft_gun") or 0) + int(v.get("eight_ft_gun") or 0),
		"K": int(v.get("six_ft_knife") or 0) + int(v.get("eight_ft_knife") or 0),
		"D": int(v.get("six_ft_display") or 0) + int(v.get("eight_ft_display") or 0),
		"N": int(v.get("six_ft_nonrelated") or 0) + int(v.get("eight_ft_nonrelated") or 0),
	}
	main_tag = max(counts, key=counts.get) if any(counts.values()) else ""

	# 👇 Add R if any rollover flag is true, regardless of per-table count
	if v.get("on_rollover") or v.get("rta") or v.get("rtn"):
		return f"{main_tag}R" if main_tag else "R"

	return main_tag



def build_lower_layout_map(template_path):
	document = Document(template_path)
	layout_map = {}
	with open("lower_layout_map_debug.txt", "w", encoding="utf-8") as debug_file:
		debug_file.write("[DEBUG] Lower Layout Map Scan (Structured)")
		for t_index, table in enumerate(document.tables):
			debug_file.write(f"[TABLE {t_index}]")
			for r_index, row in enumerate(table.rows):
				cells = [cell.text.strip() for cell in row.cells]
				debug_file.write(f"Row {r_index}: {cells}")
				for c_index in range(len(cells) - 1):
					section = cells[c_index]
					table_num = cells[c_index + 1]
					if not section or not table_num:
						continue
					if "LOWER LEVEL" in section.upper():
						continue
					key = f"{section} {table_num}"
					layout_map[key] = (t_index, r_index, c_index + 2, c_index + 3)
					debug_file.write(f"  -> {key} => Table {t_index}, Row {r_index}, VendorCol {c_index+2}, TagCol {c_index+3}")
	return layout_map






def place_vendor(document, layout_map, vendor):
	row = (vendor.get("row") or "").strip().upper()
	island = (vendor.get("island") or "").strip()
	table_numbers = parse_table_numbers(vendor.get("table_numbers", ""))
	if not row or not table_numbers:
		with open("skipped_vendors.txt", "a", encoding="utf-8") as skip_file:
			skip_file.write(f"[SKIPPED] Vendor: {vendor.get('business_name', '')} | Reason: Missing row or table numbers")
		return False

	# Determine display names for each table
	for idx, t in enumerate(table_numbers):
		if row in ["LA", "LB", "LC", "LD", "XLN"] and island:
			section = f"{island}{row}"
		else:
			section = row
		key = f"{section} {t}"

		# Determine display_name logic
		name = vendor.get("business_name", "").strip()
		if not name:
			first = vendor.get("first_name", "").strip()
			last = vendor.get("last_name", "").strip()
			name = f"{first} {last}".strip()
			if not name:
				name = "UNKNOWN"
		display_name = name if idx == 0 else "↓"

		if key in layout_map:
			t_index, r_index, vendor_col, tag_col = layout_map[key]
			table = document.tables[t_index]
			vendor_cell = table.rows[r_index].cells[vendor_col]
			tag_cell = table.rows[r_index].cells[tag_col]

			p = vendor_cell.paragraphs[0]
			p.clear()
			p.add_run(display_name)

			p2 = tag_cell.paragraphs[0]
			p2.clear()
			p2.add_run(build_tags(vendor))

			# Apply alignment
			for paragraph in vendor_cell.paragraphs:
				paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
			for paragraph in tag_cell.paragraphs:
				paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

			# Apply cell colors
			if vendor.get("electricity", False):
				set_cell_background(vendor_cell, COLOR_INDIGO)
				set_cell_background(tag_cell, COLOR_INDIGO)
			elif vendor.get("on_rollover", False) or vendor.get("rta", False) or vendor.get("rtn", False):
				set_cell_background(vendor_cell, COLOR_YELLOW)
				set_cell_background(tag_cell, COLOR_YELLOW)
			elif vendor.get("saturday_setup", False):
				set_cell_background(vendor_cell, COLOR_LIME)
				set_cell_background(tag_cell, COLOR_LIME)

			# Log clean debug output
			with open("vendor_table_mapping_debug.txt", "a", encoding="utf-8") as log_file:
				log_file.write(
					f"[INFO] Vendor: {vendor.get('business_name', '').strip() or 'N/A'} | Key: {key} | Display: {display_name} | Tag: {build_tags(vendor)}"
				)
		else:
			with open("vendor_table_mapping_debug.txt", "a", encoding="utf-8") as log_file:
				log_file.write(f"[WARN] No mapping for {vendor.get('business_name', '')} with key {key}")
	return True


def apply_vendor_color(cell, vendor):
	if vendor.get('on_rollover') or vendor.get('rta') or vendor.get('rtn'):
		set_cell_background(cell, COLOR_YELLOW)
	elif vendor.get('electricity'):
		set_cell_background(cell, COLOR_INDIGO)
	elif vendor.get('saturday_setup'):
		set_cell_background(cell, COLOR_LIME)


def log_warn(message):
	with open("vendor_table_mapping_debug.txt", "a", encoding="utf-8") as log_file:
		log_file.write(f"[WARN] {message}\n")


def log_info(message):
	with open("vendor_table_mapping_debug.txt", "a", encoding="utf-8") as log_file:
		log_file.write(f"[INFO] {message}\n")


import ctypes


def check_double_bookings(vendors, filename="double_booking_report.txt"):
	conflicts = []
	seen = {}

	for v in vendors:
		island = (v.get("island") or "").strip().upper()
		row = (v.get("row") or "").strip().upper()
		name = v.get("business_name") or f"{v.get('first_name', '')} {v.get('last_name', '')}".strip()
		tables = parse_table_numbers(v.get("table_numbers"))

		for t in tables:
			key = (island, row, t) if island else (row, t)
			if key in seen:
				conflicts.append(f"{' '.join([i for i in [island, row] if i])} Table {t} -> {seen[key]} & {name}")
			else:
				seen[key] = name

	if conflicts:
		with open(filename, "w", encoding="utf-8") as f:
			f.write("Double Booking Report\n")
			f.write("====================\n\n")
			for c in conflicts:
				f.write(c + "\n")

		conflict_msg = f"{len(conflicts)} double bookings detected! See {filename}"
		print(f"[WARNING] {conflict_msg}")
		try:
			ctypes.windll.user32.MessageBoxW(0, conflict_msg, "Double Booking Alert", 0x10)
		except Exception:
			print("[WARN] Unable to show popup alert.")
		raise SystemExit(f"Aborting due to double bookings. Check {filename} for details.")
	else:
		print("[INFO] No double bookings detected.")


def generate_layouts(vendors, template):
	document = Document("Master Layout Lower.docx")
	layout_map = build_lower_layout_map("Master Layout Lower.docx")
	open("vendor_table_mapping_debug.txt", "w").close()
	open("skipped_vendors.txt", "w").close()
	for vendor in vendors:
		place_vendor(document, layout_map, vendor)
	output_file = "Lower_Filled.docx"
	document.save(output_file)
	print(f"[INFO] Lower layout saved to {output_file}")


def main():
	parser = argparse.ArgumentParser()
	parser.add_argument("--template", required=True)
	parser.add_argument("--pg_host", required=True)
	parser.add_argument("--pg_port", type=int, required=True)
	parser.add_argument("--pg_dbname", required=True)
	parser.add_argument("--pg_user", required=True)
	parser.add_argument("--pg_password", required=True)
	args = parser.parse_args()

	vendors = load_vendors_from_postgres(
		args.pg_host, args.pg_port, args.pg_dbname, args.pg_user, args.pg_password
	)
	generate_layouts(vendors, args.template)


if __name__ == "__main__":
	main()
